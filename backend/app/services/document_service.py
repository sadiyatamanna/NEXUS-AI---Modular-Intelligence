import os
import re
from typing import List, Dict, Any, Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument

class DocumentParser:
    """Extracts clean text and structures chunks from PDF, DOCX, Markdown, and TXT files."""

    @staticmethod
    def extract_text(file_path: str, mime_type: str) -> List[Tuple[int, str]]:
        """Returns a list of tuples: (page_or_section_number, text_content)."""
        pages = []
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf" or "pdf" in mime_type:
            try:
                reader = PdfReader(file_path, strict=False)
                for idx, page in enumerate(reader.pages):
                    try:
                        text = page.extract_text() or ""
                    except Exception:
                        text = ""
                    clean_text = DocumentParser._clean_text(text)
                    if clean_text:
                        pages.append((idx + 1, clean_text))
            except Exception as e:
                raise ValueError(f"Failed to parse PDF document: {e}")

        elif ext == ".docx" or "word" in mime_type or "officedocument" in mime_type:
            try:
                doc = DocxDocument(file_path)
                text_parts = []
                # 1. Extract paragraphs
                for p in doc.paragraphs:
                    p_txt = p.text.strip()
                    if p_txt:
                        text_parts.append(p_txt)
                # 2. Extract tables (vital for structured reports & review forms)
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            text_parts.append(" | ".join(row_cells))
                
                full_text = "\n\n".join(text_parts)
                clean_text = DocumentParser._clean_text(full_text)
                if clean_text:
                    pages.append((1, clean_text))
            except Exception as e:
                raise ValueError(f"Failed to parse DOCX document: {e}")

        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"] or "image" in mime_type:
            try:
                from PIL import Image
                with Image.open(file_path) as img:
                    filename = os.path.basename(file_path)
                    desc = (
                        f"Image Document: {filename}\n"
                        f"Format: {img.format}, Dimensions: {img.width}x{img.height} pixels, Mode: {img.mode}.\n"
                        "Visual document indexed into knowledge base."
                    )
                    pages.append((1, desc))
            except Exception:
                pages.append((1, f"Image document: {os.path.basename(file_path)}"))

        else:
            # Fallback for plain text, markdown, CSV, JSON, code, etc.
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    content = f.read()
                    clean_text = DocumentParser._clean_text(content)
                    if clean_text:
                        pages.append((1, clean_text))
            except Exception as e:
                raise ValueError(f"Failed to read file: {e}")

        return pages

    @staticmethod
    def _clean_text(text: str) -> str:
        # Normalize whitespace
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n\s*\n+", "\n\n", text)
        return text.strip()

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 600, overlap: int = 100) -> List[str]:
        """Split text into overlapping semantic passages."""
        words = text.split()
        if not words:
            return []

        chunks = []
        start = 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            chunk = " ".join(words[start:end])
            chunks.append(chunk)
            if end == len(words):
                break
            start += chunk_size - overlap
        return chunks
